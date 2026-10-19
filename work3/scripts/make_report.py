from pathlib import Path
import argparse, json, re, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
def main():
    parser=argparse.ArgumentParser(description="Сформировать отчёт по практической работе № 3.")
    parser.add_argument("--results",type=Path,default=ROOT/"results")
    parser.add_argument("--output",type=Path,default=ROOT/"results"/"report.docx")
    args=parser.parse_args()
    result_dir=args.results
    summary=json.loads((result_dir/"summary.json").read_text(encoding="utf-8"))
    tests=ET.parse(result_dir/"tests.xml").getroot()
    total=int(tests.get("tests","0")); failures=int(tests.get("failures","0"))+int(tests.get("errors","0"))
    cases=tests.findall(".//testcase")
    parameterized=[c for c in cases if "All343Shapes" in c.get("classname","")]
    if total==0 or failures or len(parameterized)!=343:
        raise SystemExit("Нельзя сформировать итоговый отчёт: tests.xml должен содержать 343 параметризованных случая без ошибок.")
    for name in ("complexity.png","speedup.png"):
        if not (result_dir/name).is_file(): raise SystemExit(f"Не найден график {name}; сначала запустите scripts/plot.py.")
    rows=summary["rows"]; ctx=summary.get("context",{})
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Cm(2); sec.bottom_margin=Cm(2); sec.left_margin=Cm(3); sec.right_margin=Cm(1.5)
    normal=doc.styles["Normal"]; normal.font.name="Times New Roman"; normal.font.size=Pt(14)
    normal.paragraph_format.line_spacing=1.5; normal.paragraph_format.space_after=Pt(5)
    for name in ("Title","Heading 1","Heading 2"):
        doc.styles[name].font.name="Times New Roman"
    doc.styles["Heading 1"].font.size=Pt(14); doc.styles["Heading 1"].font.bold=True
    doc.styles["Heading 2"].font.size=Pt(14); doc.styles["Heading 2"].font.bold=True
    def para(text="",bold=False,center=False,indent=True):
        p=doc.add_paragraph()
        if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if indent and not center:p.paragraph_format.first_line_indent=Cm(1.25)
        r=p.add_run(text);r.bold=bold
        return p
    def heading(text,level=1): doc.add_heading(text,level=level)
    def code(text):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.right_indent=Cm(.3)
        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.0
        r=p.add_run(text);r.font.name="Consolas";r.font.size=Pt(8)
    def table(headers,data):
        t=doc.add_table(rows=1,cols=len(headers));t.style="Table Grid";t.alignment=WD_TABLE_ALIGNMENT.CENTER
        for cell,value in zip(t.rows[0].cells,headers): cell.text=str(value)
        for row in data:
            cells=t.add_row().cells
            for cell,value in zip(cells,row):cell.text=str(value)
        for row in t.rows:
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0)
                    for r in p.runs:r.font.name="Times New Roman";r.font.size=Pt(10)
        return t
    para("Министерство науки и высшего образования РФ",center=True)
    para("Федеральное государственное автономное образовательное учреждение высшего образования",center=True)
    para("«СИБИРСКИЙ ФЕДЕРАЛЬНЫЙ УНИВЕРСИТЕТ»",bold=True,center=True)
    para("Институт космических и информационных технологий",center=True)
    para("Кафедра «Вычислительная техника»",center=True)
    for _ in range(4):doc.add_paragraph()
    para("ОТЧЁТ О ПРАКТИЧЕСКОЙ РАБОТЕ № 3",bold=True,center=True)
    para("Умножение матриц с использованием разделяемой памяти CUDA",center=True)
    for _ in range(3):doc.add_paragraph()
    para("Преподаватель                                                     С.А. Тарасов",indent=False)
    para("Студент           КИ23-06Б 032320481                         А.А. Платонов",indent=False)
    para("Студент           КИ23-06Б 032321113                         Н.С. Бученик",indent=False)
    for _ in range(4):doc.add_paragraph()
    para("Красноярск 2026",center=True)
    doc.add_page_break()
    heading("ВВЕДЕНИЕ")
    para("Цель работы — закрепить навыки программирования CUDA и изучить применение разделяемой памяти при умножении матриц. Требуется реализовать kernel_matmul_shmem с передачей MatrixView по значению, обеспечить корректность результата, сравнить производительность с наивным ядром и проанализировать вычислительную сложность.")
    para("Реализованы классы Data, MatrixView и Matrix. Оператор operator* использует стратегию SharedMemory; наивная реализация сохранена для контрольного сравнения. В обоих вариантах одна нить вычисляет один элемент матрицы результата, а блок содержит 16 × 16 нитей.")
    para("Тестирование охватывает все 343 сочетания размеров m, k и n из множества {1, 2, 3, 127, 128, 129, 512}. Производительность измеряется для квадратных матриц порядков 16, 32, 64, 128, 256, 512 и 1024.")
    para("Исходный код, протоколы тестирования, графики и скрипты отчёта размещены в репозитории:",indent=False)
    para("https://github.com/NikitaBuchenik/hsys/tree/main/work3",indent=False)
    doc.add_page_break()
    heading("1 Разработка программного решения")
    heading("1.1 Владение памятью: Data",2)
    para("Data владеет массивом в глобальной памяти GPU. Конструктор выделяет память с помощью cudaMalloc, а деструктор освобождает её через cudaFree. Копирование выполняет глубокую копию устройства в устройство, перемещение передаёт указатель. Проверяется переполнение размера; ошибки CUDA обрабатываются через CUDA_CHECK.")
    code('explicit Data(std::size_t size) : size_(size) {\n    if (size > max_size / sizeof(T)) throw std::length_error("Data size overflow");\n    if (size_) CUDA_CHECK(cudaMalloc(reinterpret_cast<void**>(&ptr_), size_ * sizeof(T)));\n}')
    heading("1.2 Представление MatrixView",2)
    para("MatrixView хранит невладеющий указатель и размеры матрицы. Элемент (i, j) расположен по индексу i · ncols + j, то есть используется row-major. Объект не управляет памятью и передаётся в CUDA-ядро по значению. Для входных матриц используется MatrixView<const T>, запрещающий изменение исходных элементов через представление.")
    code('template<class T> class MatrixView {\n    T* ptr_; std::size_t rows_, cols_;\npublic:\n    __host__ __device__ T& operator()(std::size_t r, std::size_t c) const {\n        return ptr_[r * cols_ + c];\n    }\n};')
    heading("1.3 Оператор умножения и стратегия",2)
    para("Перед вычислением проверяется условие A.ncols() == B.nrows(). Создаётся матрица результата; при ненулевом размере запускается выбранное ядро, после чего cudaDeviceSynchronize ожидает завершения вычисления и позволяет обнаружить ошибки выполнения.")
    code('enum class MatmulStrategy { Naive, SharedMemory };\ntemplate<class T> Matrix<T> operator*(const Matrix<T>& a, const Matrix<T>& b) {\n    return multiply(a, b, MatmulStrategy::SharedMemory);\n}')
    para("Выбор стратегии отделяет интерфейс умножения от реализации алгоритма. multiply_naive оставлен для воспроизводимого сравнения с исходным подходом.")
    heading("1.4 CUDA-ядро с разделяемой памятью",2)
    para("Блок из 256 нитей обрабатывает плитку результата 16 × 16. На каждой итерации цикла по внутреннему измерению потоки загружают фрагменты A и B в две плитки shared memory. Для выходящих за границы элементов записываются нули. После загрузки вызывается __syncthreads(), затем каждая нить накапливает частичную сумму; повторная синхронизация защищает плитки перед следующей загрузкой.")
    code('tile_a[ty][tx] = valid_a ? a(row, ak) : T{0};\ntile_b[ty][tx] = valid_b ? b(bk, col) : T{0};\n__syncthreads();\n#pragma unroll\nfor (unsigned j = 0; j < TILE; ++j)\n    sum += tile_a[ty][j] * tile_b[j][tx];\n__syncthreads();')
    para("Обработка границ позволяет вычислять размеры, не кратные 16. Для квадратных матриц обе реализации имеют арифметическую сложность Θ(n³); использование плиток не меняет асимптотику, но уменьшает повторные чтения данных из глобальной памяти.")
    doc.add_page_break()
    heading("2 Модульное тестирование")
    para("Google Test строит декартово произведение трёх наборов размеров через Combine. Эталонный результат вычисляется с помощью Eigen::MatrixXf. Входные матрицы хранятся в row-major, что согласуется с раскладкой MatrixView.")
    code('Eigen::MatrixXf expected = Eigen::MatrixXf(ar) * Eigen::MatrixXf(br);\nMatrix<float> a(m, k), b(k, n);\na.data().copy_from_host(ar.data());\nb.data().copy_from_host(br.data());\nauto c = a * b;\nEXPECT_TRUE(actual.isApprox(expected, 1e-5f));\nEXPECT_LE((actual - expected).cwiseAbs().maxCoeff(), 1e-5f);')
    para(f"По сохранённому протоколу tests.xml выполнено {total} тестов; ошибок и падений — {failures}. В параметризованном наборе All343Shapes присутствуют все 343 комбинации размеров. Дополнительно проверяются некратные плитке размеры, несовместимые размеры, нулевое внутреннее измерение и совпадение наивной и оптимизированной реализаций.")
    para("isApprox проверяет относительную близость норм, поэтому отдельно контролируется максимальная абсолютная погрешность не более 10⁻⁵.")
    doc.add_page_break()
    heading("3 Методика измерения производительности")
    para("Измерения выполнены Google Benchmark в Release-сборке. Для каждого размера проведены пять повторений; в таблице приведены медианные значения. Матрицы создаются и копируются на GPU до измеряемого интервала. CUDA Events измеряют выполнение ядра, поэтому выделение памяти результата и освобождение памяти не включены в измеряемый интервал.")
    para("Перед измерениями выполняются прогревочные запуски. События CUDA синхронизируются перед считыванием elapsed time. Измерения отражают время вычислительной части и не включают передачу исходных данных между CPU и GPU; полное время приложения будет больше.")
    para("Среда эксперимента: NVIDIA GeForce RTX 2060, 6 ГиБ памяти; CUDA Compiler 13.0.88; NVIDIA Driver 595.79; CUDA Version, указанная nvidia-smi, — 13.2. Google Benchmark v1.9.1, Release. В метаданных Google Benchmark указаны 12 логических процессоров и частота 3793 МГц.")
    doc.add_page_break()
    heading("4 Результаты вычислительного эксперимента")
    para("В таблице приведены медианы времени для двух CUDA-ядер и отношение времени наивного ядра к времени реализации с разделяемой памятью. Ускорение больше единицы означает преимущество shared-memory реализации.")
    data=[]
    for r in rows:
        data.append([r["n"],f'{r["naive_seconds"]*1000:.3f}'.replace(".",","),f'{r["shared_memory_seconds"]*1000:.3f}'.replace(".",","),f'{r["speedup_naive_over_shmem"]:.3f}'.replace(".",",")])
    table(["n","Наивное ядро, мс","Shared memory, мс","Ускорение"],data)
    para("Исходные значения измерений сохранены в results/benchmarks.json и results/measurements.csv. Графики построены скриптом plot.py на основании измеренных данных, а не на основе теоретических оценок.")
    doc.add_picture(str(result_dir/"complexity.png"),width=Cm(15))
    para("Рисунок 1 — Время вычислений в зависимости от размера матрицы",center=True,indent=False)
    doc.add_page_break()
    heading("4.1 Ускорение новой реализации")
    doc.add_picture(str(result_dir/"speedup.png"),width=Cm(15))
    para("Рисунок 2 — Отношение времени наивного ядра к времени shared-memory ядра",center=True,indent=False)
    para("Для n = 16 и n = 32 ускорение составляет соответственно 0,952 и 0,823: на этих размерах накладные расходы загрузки плиток и синхронизации не окупаются. Для n = 64 и n = 128 выигрыш небольшой — около 1,04 и 1,03. При n = 256 ускорение достигает 1,216, при n = 512 — 1,263, а при n = 1024 — 1,292. Наибольший измеренный выигрыш в этом эксперименте составляет примерно 29 %.")
    heading("4.2 Сопоставление с теоретической сложностью")
    para("Теоретическая сложность умножения квадратных матриц равна Θ(n³). При удвоении n время для больших размеров растёт приблизительно в восемь раз. На интервале 256 → 512 отношение времени составляет около 6,69 для shared memory и 6,95 для наивного ядра; на интервале 512 → 1024 — около 7,84 и 8,02 соответственно. Эти значения приближаются к ожидаемому коэффициенту 8.")
    para("На малых размерах экспериментальная кривая не обязана быть монотонной: стоимость запуска ядра, таймера, планирования блоков и синхронизации сопоставима с полезной работой. Поэтому небольшие отклонения от кубической модели ожидаемы. Ускорение зависит от архитектуры GPU и размеров матриц; результаты одного устройства не являются универсальной гарантией.")
    heading("5 Вывод")
    para("Реализовано умножение матриц с применением shared memory и передачи MatrixView по значению. Протокол содержит 347 успешно завершённых тестов, включая все 343 заданные комбинации размеров. По измерениям на NVIDIA GeForce RTX 2060 разделяемая память не даёт преимущества для самых маленьких матриц, но обеспечивает ускорение на крупных размерах до 1,292 раза при n = 1024. Экспериментальные показатели для крупных матриц согласуются с теоретической сложностью Θ(n³).")
    para("Полный проект и воспроизводимые данные эксперимента: https://github.com/NikitaBuchenik/hsys/tree/main/work3",indent=False)
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    print(f"Report saved to {args.output}; validated {total} tests and 343 parameterized shapes.")
if __name__=="__main__":main()

# DOCX report is generated in CI from the committed test and benchmark artifacts.
